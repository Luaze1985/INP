import re
import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders(table, color="CBD5E0"):
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
        <w:left w:val="none"/>
        <w:bottom w:val="single" w:sz="6" w:space="0" w:color="{color}"/>
        <w:right w:val="none"/>
        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color}"/>
        <w:insideV w:val="none"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))

def add_inline_formatted_text(paragraph, text, color=None):
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith('*') and token.endswith('*'):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = color or RGBColor(0x80, 0x20, 0x20)
        else:
            cleaned = re.sub(r'\[([^\]]+)\]\([^)]+\)', r'\1', token)
            run = paragraph.add_run(cleaned)
        if color and not (token.startswith('`') and token.endswith('`')):
            run.font.color.rgb = color


def add_formatted_text(paragraph, text):
    parts = re.split(r'(<span style="color:red">.*?</span>)', text)
    for part in parts:
        if not part:
            continue
        if part.startswith('<span style="color:red">') and part.endswith('</span>'):
            add_inline_formatted_text(
                paragraph,
                part[len('<span style="color:red">'):-len('</span>')],
                RGBColor(0xFF, 0x00, 0x00),
            )
        else:
            add_inline_formatted_text(paragraph, part)

def convert_md_to_docx(md_path, docx_path):
    doc = docx.Document()

    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Segoe UI'
    font.size = Pt(10.5)
    font.color.rgb = RGBColor(0x2D, 0x37, 0x48)

    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    i = 0
    table_lines = []

    def flush_table(lines):
        if not lines:
            return
        rows_data = []
        for line in lines:
            if '|' in line:
                cells = [c.strip() for c in line.strip().strip('|').split('|')]
                if all(re.match(r'^:?-+:?$', c) for c in cells):
                    continue
                rows_data.append(cells)

        if not rows_data:
            return

        num_rows = len(rows_data)
        num_cols = max(len(r) for r in rows_data)
        table = doc.add_table(rows=num_rows, cols=num_cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, "CBD5E0")

        for r_idx, row in enumerate(rows_data):
            tr = table.rows[r_idx]
            is_header = (r_idx == 0)
            for c_idx, cell_text in enumerate(row):
                if c_idx < len(tr.cells):
                    cell = tr.cells[c_idx]
                    set_cell_margins(cell, top=120, bottom=120, left=150, right=150)
                    if is_header:
                        set_cell_background(cell, "1B365D")
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(4)
                        p.paragraph_format.space_after = Pt(4)
                        run = p.add_run(cell_text)
                        run.bold = True
                        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                        run.font.size = Pt(10)
                    else:
                        if r_idx % 2 == 1:
                            set_cell_background(cell, "F7FAFC")
                        else:
                            set_cell_background(cell, "FFFFFF")
                        p = cell.paragraphs[0]
                        p.paragraph_format.space_before = Pt(3)
                        p.paragraph_format.space_after = Pt(3)
                        add_formatted_text(p, cell_text)

        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(0)
        p_space.paragraph_format.space_after = Pt(6)

    while i < len(lines):
        line = lines[i].rstrip('\r\n')

        if line.startswith('|'):
            table_lines.append(line)
            i += 1
            continue
        elif table_lines:
            flush_table(table_lines)
            table_lines = []

        if not line.strip():
            i += 1
            continue

        if line.strip() in ['---', '***', '___']:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:bottom w:val="single" w:sz="6" w:space="1" w:color="CBD5E0"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            i += 1
            continue

        if line.startswith('# '):
            h = doc.add_heading(level=1)
            h.paragraph_format.space_before = Pt(16)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(line[2:].strip())
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
            run.font.size = Pt(18)
            run.bold = True
        elif line.startswith('## '):
            h = doc.add_heading(level=2)
            h.paragraph_format.space_before = Pt(14)
            h.paragraph_format.space_after = Pt(4)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(line[3:].strip())
            run.font.color.rgb = RGBColor(0x2C, 0x52, 0x82)
            run.font.size = Pt(14)
            run.bold = True
        elif line.startswith('### '):
            h = doc.add_heading(level=3)
            h.paragraph_format.space_before = Pt(10)
            h.paragraph_format.space_after = Pt(3)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(line[4:].strip())
            run.font.color.rgb = RGBColor(0x2B, 0x6C, 0xB0)
            run.font.size = Pt(12)
            run.bold = True
        elif line.startswith('#### '):
            h = doc.add_heading(level=4)
            h.paragraph_format.space_before = Pt(8)
            h.paragraph_format.space_after = Pt(2)
            h.paragraph_format.keep_with_next = True
            run = h.add_run(line[5:].strip())
            run.font.color.rgb = RGBColor(0x4A, 0x55, 0x68)
            run.font.size = Pt(11)
            run.bold = True
        elif line.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.left_indent = Inches(0.3)
            pBdr = parse_xml(f'<w:pBdr {nsdecls("w")}><w:left w:val="single" w:sz="18" w:space="12" w:color="2B6CB0"/></w:pBdr>')
            p._p.get_or_add_pPr().append(pBdr)
            add_formatted_text(p, line[2:].strip())
        elif line.startswith('- ') or line.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            add_formatted_text(p, line[2:].strip())
        elif re.match(r'^\d+\.\s+', line):
            content = re.sub(r'^\d+\.\s+', '', line)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            add_formatted_text(p, content)
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            add_formatted_text(p, line.strip())

        i += 1

    if table_lines:
        flush_table(table_lines)

    os.makedirs(os.path.dirname(docx_path), exist_ok=True)
    doc.save(docx_path)
    print(f"Successfully generated: {docx_path}")

if __name__ == '__main__':
    import sys
    if len(sys.argv) >= 3:
        convert_md_to_docx(sys.argv[1], sys.argv[2])
    else:
        print("Usage: python md_to_docx.py <input.md> <output.docx>")
